"""Teacher-only Excel export with rubric levels and individual feedback."""
import io
import json
from openpyxl import Workbook
from openpyxl.styles import Alignment, Font, PatternFill
from assessment import EXAMS, evaluate
from rubrics import RUBRICS, band_feedback, rubric_rows

def export_workbook(submissions):
    wb = Workbook()
    summary = wb.active
    summary.title = '학생별 성적'
    summary.append(['평가', '학생 아이디', '이름', '자동 채점 합계', '미확정 항목 수', '교사 확정 합계', '만점', '제출 시각(UTC)'])
    details = wb.create_sheet('문항별 평가')
    details.append(['평가', '학생 아이디', '이름', '문항', '답안', '자동 점수', '교사 확정 점수', '만점', '점수 구간 기준', '점수별 피드백', '역사적 사실 피드백', '근거 자료', '교사 추가 피드백'])
    for row in submissions:
        exam = row['exam']
        answers = json.loads(row['answers'])
        result = evaluate(exam, answers)
        review = json.loads(row['review']) if row.get('review') else {}
        keys = list(RUBRICS[exam])
        pending = sum(result[k]['score'] is None for k in keys)
        provisional = sum(result[k]['score'] for k in keys) if not pending else None
        final = sum(review[k]['score'] for k in keys) if all(k in review for k in keys) else None
        summary.append([EXAMS[exam]['title'], row['student'], row['name'], provisional, pending, final, 20, row['submitted']])
        for key, number, _, points, _ in EXAMS[exam]['questions']:
            if not points:
                continue
            item = result[key]
            confirmed = review.get(key, {})
            score = confirmed.get('score', item['score'])
            levels = RUBRICS[exam][key][1]
            descriptor = levels[score][0] if score in levels else item['rubric_descriptor']
            feedback = band_feedback(exam, key, score) if score in levels else item['score_feedback']
            answer = (answers.get('event', '') + '\n' if key == 'production' else '') + answers.get(key, '')
            details.append([EXAMS[exam]['title'], row['student'], row['name'], number, answer, item['score'], confirmed.get('score'), points, descriptor, feedback, item['feedback'], item['rubric_source']['file'] + ' · PDF 1쪽', confirmed.get('feedback', '')])
    for exam in RUBRICS:
        sheet = wb.create_sheet('문항1 루브릭' if exam == 'industry' else '문항2 루브릭')
        sheet.append(['평가 항목', '점수', '채점 기준', '점수별 피드백'])
        for record in rubric_rows(exam):
            sheet.append(list(record.values()))
    note = wb.create_sheet('운영 안내')
    note.append(['항목', '설명'])
    note.append(['자동 점수', '규칙 기반 자동 점수입니다. 문맥 확인이 필요한 항목은 공란입니다. 교사가 최종 확정합니다.'])
    note.append(['1번 문항', '1-①과 1-② 답안을 합쳐 4점으로 채점합니다. 별도 가산하지 않습니다.'])
    note.append(['최저 점수', '원본 표의 최저 점수는 항목당 1점입니다. 0점을 새로 만들지 않았습니다.'])
    note.append(['표기 차이', '문항2 마지막 항목 원본은 영국 횡단정책으로 표기. 피드백 내용.pdf 4~5쪽에 따라 종단정책 적용.'])
    for sheet in wb:
        sheet.freeze_panes = 'A2'
        sheet.auto_filter.ref = sheet.dimensions
        for cell in sheet[1]:
            cell.fill = PatternFill('solid', fgColor='2E6354')
            cell.font = Font(color='FFFFFF', bold=True)
        for row in sheet.iter_rows(min_row=2):
            for cell in row:
                # All supplied strings remain text, including formula-like answers.
                if isinstance(cell.value, str):
                    cell.data_type = 's'
                cell.alignment = Alignment(vertical='top', wrap_text=True)
        for column in sheet.columns:
            length = max(len(str(c.value or '')) for c in column)
            sheet.column_dimensions[column[0].column_letter].width = min(65, max(14, length + 2))
    output = io.BytesIO()
    wb.save(output)
    return output.getvalue()
