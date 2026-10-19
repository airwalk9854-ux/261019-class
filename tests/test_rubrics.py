import io
import json
import unittest
from openpyxl import load_workbook
from assessment import evaluate
from gradebook import export_workbook
from rubrics import RUBRICS, bands, score_answer, band_feedback


class RubricTests(unittest.TestCase):
    def test_original_discrete_bands(self):
        expected = {'industry': {'production': [4, 3, 2, 1], 'positive': [4, 3, 2, 1], 'negative': [4, 3, 2, 1], 'critique': [8, 6, 4, 3, 2, 1]},
                    'imperialism': {'production': [4, 3, 2, 1], 'rhodes': [6, 4, 1], 'france': [4, 3, 2, 1], 'africa': [6, 4, 1]}}
        for exam, criteria in expected.items():
            self.assertEqual(sum(max(bands(exam, key)) for key in criteria), 20)
            for key, values in criteria.items():
                self.assertEqual(list(bands(exam, key)), values)
                feedback = [band_feedback(exam, key, value) for value in values]
                self.assertEqual(len(set(feedback)), len(values))
                self.assertNotIn('산업혁명', ''.join(feedback))
                self.assertNotIn('횡단정책', ''.join(feedback))
                self.assertNotIn('종단정책', ''.join(feedback))
                self.assertNotIn('파쇼다', ''.join(feedback))

    def test_all_score_bands_with_representative_answers(self):
        cases = [
            ('industry', 'production', 4, '산업혁명으로 가내 수공업에서 공장제 기계공업으로 변하여 면직물 생산량이 증가하였다.'),
            ('industry', 'production', 3, '가내 수공업에서 공장제 기계공업으로 변하여 면직물 생산량이 증가하였다.'),
            ('industry', 'production', 2, '산업혁명으로 가내 수공업에서 공장제 기계공업으로 변하였다.'),
            ('industry', 'production', 1, ''),
            ('industry', 'positive', 4, '만국 박람회에서 발전된 기계를 전시하여 산업과 국력을 과시했다.'),
            ('industry', 'positive', 4, '철도로 생활 범위와 여가 활동이 확대되고 시간표를 통해 시간 개념이 생겼다.'),
            ('industry', 'positive', 3, '박람회를 통해 산업혁명의 성과를 알렸다.'),
            ('industry', 'positive', 3, '철도로 생활 범위가 넓어지고 여가를 즐겼다.'),
            ('industry', 'positive', 2, '영국 만국 박람회가 긍정적인 사례다.'),
            ('industry', 'positive', 1, ''),
            ('industry', 'negative', 4, '인터뷰에 따르면 아동에 대한 노동 착취가 심각하였다.'),
            ('industry', 'negative', 4, '환경 오염으로 위생이 열악한 도시 빈민가에 콜레라가 유행했다.'),
            ('industry', 'negative', 3, '아동이 힘들게 노동했다.'),
            ('industry', 'negative', 3, '환경 오염으로 전염병이 유행했다.'),
            ('industry', 'negative', 2, '아동 노동 문제이다.'),
            ('industry', 'negative', 1, ''),
            ('industry', 'critique', 8, '박람회는 산업과 국력의 발전을 보여주지만 식민지 전시장은 식민지 주민의 착취를 보여 준다.'),
            ('industry', 'critique', 6, '박람회는 산업 발전을 보여주지만 러다이트 운동 같은 노동 문제가 발생했다.'),
            ('industry', 'critique', 4, '박람회는 산업 발전을 보여주지만 노동 문제와 빈부 격차가 발생했다.'),
            ('industry', 'critique', 3, '식민지 사람들을 착취하였다.'),
            ('industry', 'critique', 2, '박람회는 산업 발전을 과시하였다.'),
            ('industry', 'critique', 1, ''),
            ('imperialism', 'rhodes', 6, '실업 문제 해결과 시장 확보를 위해 제국주의를 추진했고 많은 식민지를 확보하였다.'),
            ('imperialism', 'rhodes', 4, '실업 문제와 시장 확보가 정책의 원인이다.'),
            ('imperialism', 'rhodes', 4, '영국은 많은 식민지를 확보하였다.'),
            ('imperialism', 'rhodes', 1, ''),
            ('imperialism', 'france', 4, '우리 프랑스의 횡단정책이 영국과 파쇼다 사건에서 충돌하여 위협을 느꼈다.'),
            ('imperialism', 'france', 3, '프랑스는 횡단정책을 추진했다.'),
            ('imperialism', 'france', 3, '파쇼다 사건이 일어났다.'),
            ('imperialism', 'france', 2, '프랑스와 영국은 아프리카 식민지를 두고 경쟁했다.'),
            ('imperialism', 'france', 1, ''),
            ('imperialism', 'africa', 6, '영국의 종단정책으로 우리 땅을 빼앗겨 침략에 분노한다.'),
            ('imperialism', 'africa', 4, '제국주의는 식민지 지배를 확대했다.'),
            ('imperialism', 'africa', 1, ''),
        ]
        for exam, key, expected, text in cases:
            with self.subTest(exam=exam, key=key, expected=expected, text=text):
                actual = score_answer(exam, key, {key: text})
                self.assertEqual(actual['score'], expected)
                self.assertEqual(actual['score_feedback'], band_feedback(exam, key, expected))

    def test_joint_first_question_and_abstention(self):
        result = evaluate('industry', {'event': '산업혁명', 'production': '가내 수공업에서 공장제 기계공업으로 바뀌어 면직물 생산량이 증가했다.'})
        self.assertEqual(result['event']['score'], 0)
        self.assertEqual(result['production']['score'], 4)
        self.assertIsNone(score_answer('industry', 'production', {'production': '산업혁명으로 공장제 생산이 시작했으나 면직물은 증가하지 않았다.'})['score'])
        self.assertIsNone(score_answer('imperialism', 'africa', {'africa': '영국의 횡단정책으로 우리 땅이 침략당했다.'})['score'])
        self.assertIsNone(score_answer('industry', 'production', {'production': '면직물 생산량이 증가했다.'})['score'])
        self.assertEqual(score_answer('imperialism', 'rhodes', {'rhodes': '실업 해결과 새 식민지를 확보하기 위해 정책을 추진했다.'})['score'], 4)

    def test_excel_export_preserves_scores_and_untrusted_answer_as_text(self):
        row = {'exam': 'industry', 'student': '=SUM(1,1)', 'name': '학생', 'answers': json.dumps({'event': '산업혁명', 'positive': '=1+1'}), 'review': None, 'submitted': '2026-10-19'}
        data = export_workbook([row])
        workbook = load_workbook(io.BytesIO(data))
        self.assertEqual(workbook.sheetnames, ['학생별 성적', '문항별 평가', '문항1 루브릭', '문항2 루브릭', '운영 안내'])
        self.assertEqual(workbook['학생별 성적']['B2'].data_type, 's')
        self.assertEqual(workbook['문항별 평가']['E3'].value, '=1+1')
        self.assertEqual(workbook['문항별 평가']['E3'].data_type, 's')
        for record in workbook['문항별 평가'].iter_rows(min_row=2, values_only=True):
            self.assertEqual(record[11], '논술형 문항1 답안.pdf · PDF 1쪽')

    def test_store_rejects_out_of_band_score(self):
        import tempfile
        from pathlib import Path
        from storage import Store
        with tempfile.TemporaryDirectory() as folder:
            store = Store(Path(folder) / 'test.db')
            store.submit('imperialism', 's', '학생', {}, {})
            with self.assertRaises(ValueError):
                store.review('imperialism', 's', {'rhodes': {'score': 5}})
